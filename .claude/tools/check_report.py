"""Check the generated report against the teacher's formatting rules.

Run after build_report.py. Prints violations, or says everything is clean.
Rules come from Notes/Оформление отчётов.md.
"""

import io
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

REPORTS = {
    "1": Path("Task 1st Deep Interview Guide") / "Solution" / "Отчёт по ПР №1.docx",
    "2": Path("Task 2nd Deep Interview Case") / "Solution" / "Отчёт по ПР №2.docx",
}

# Номер работы первым аргументом, по умолчанию первая.
NUMBER = sys.argv[1] if len(sys.argv) > 1 else "1"
if NUMBER not in REPORTS:
    raise SystemExit(f"нет отчёта для работы {NUMBER}, доступны: {', '.join(REPORTS)}")
PATH = Path(__file__).resolve().parents[2] / REPORTS[NUMBER]
TEXT_WIDTH_CM = 16.5

doc = Document(PATH)
problems: list[str] = []

# --- page setup
sec = doc.sections[0]
if abs(sec.page_width.cm - 21.0) > 0.05 or abs(sec.page_height.cm - 29.7) > 0.05:
    problems.append(
        f"формат страницы {sec.page_width.cm:.1f}x{sec.page_height.cm:.1f} см вместо A4"
    )
margins = {
    "левое": (sec.left_margin.cm, 3.0),
    "правое": (sec.right_margin.cm, 1.5),
    "верхнее": (sec.top_margin.cm, 2.0),
    "нижнее": (sec.bottom_margin.cm, 2.0),
}
for name, (got, want) in margins.items():
    if abs(got - want) > 0.05:
        problems.append(f"поле {name}: {got:.2f} см вместо {want} см")

normal = doc.styles["Normal"]
if normal.font.name != "Times New Roman":
    problems.append(f"шрифт стиля Normal: {normal.font.name}")
if normal.font.size.pt != 14:
    problems.append(f"кегль стиля Normal: {normal.font.size.pt}")
pf = normal.paragraph_format
if pf.line_spacing != 1.0:
    problems.append(f"межстрочный интервал: {pf.line_spacing}")
if abs(pf.first_line_indent.cm - 1.5) > 0.02:
    problems.append(f"отступ первой строки: {pf.first_line_indent.cm:.2f} см")

# --- runs: no bold, italic, underline; right font everywhere
def scan_runs(paragraphs, where: str) -> None:
    for p in paragraphs:
        for run in p.runs:
            if not run.text.strip():
                continue
            label = run.text.strip()[:40]
            if run.bold:
                problems.append(f"жирный ({where}): {label}")
            if run.italic:
                problems.append(f"курсив ({where}): {label}")
            if run.underline:
                problems.append(f"подчёркивание ({where}): {label}")
            if run.font.name not in (None, "Times New Roman"):
                problems.append(f"шрифт {run.font.name} ({where}): {label}")
            if run.font.size is not None and run.font.size.pt != 14:
                problems.append(f"кегль {run.font.size.pt} ({where}): {label}")


scan_runs(doc.paragraphs, "текст")
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            scan_runs(cell.paragraphs, "таблица")

# --- em dash must not appear anywhere
paras = [p.text for p in doc.paragraphs]
for text in paras:
    if "—" in text:
        problems.append(f"длинное тире: {text.strip()[:60]}")

# --- table widths equal the paragraph width, captions and references
captions = [t for t in paras if t.strip().startswith("Таблица ")]
if len(captions) != len(doc.tables):
    problems.append(
        f"подписей к таблицам {len(captions)}, а таблиц {len(doc.tables)}"
    )

body = "\n".join(paras)
for i in range(1, len(doc.tables) + 1):
    if not re.search(rf"в таблице {i}\b", body):
        problems.append(f"в тексте нет ссылки на таблицу {i}")

for i, t in enumerate(doc.tables, start=1):
    total = sum(c.width.cm for c in t.rows[0].cells)
    if abs(total - TEXT_WIDTH_CM) > 0.15:
        problems.append(
            f"ширина таблицы {i}: {total:.2f} см вместо {TEXT_WIDTH_CM} см"
        )

# --- headings numbered without a dot
for text in paras:
    if re.match(r"^\d+\.\s+[А-ЯЁ]", text.strip()) and len(text.strip()) < 60:
        problems.append(f"точка после номера раздела: {text.strip()[:50]}")

# --- summary
questions = [p for p in paras if re.match(r"^\d+ .*\?", p.strip())]
numbered_items = [p for p in paras if re.match(r"^\d+ [А-ЯЁ]", p.strip())]
bullet_items = [p for p in paras if p.strip().startswith("– ")]

print(f"файл: {PATH.name}, {PATH.stat().st_size} байт")
print(f"абзацев: {len(paras)}, таблиц: {len(doc.tables)}")
print(f"пунктов гайда: {len(numbered_items)} (из них вопросов: {len(questions)})")
print(f"маркированных пунктов: {len(bullet_items)}")
print(f"подписей к таблицам: {len(captions)}")
print()

if problems:
    print(f"НАРУШЕНИЯ ({len(problems)}):")
    for problem in problems:
        print(f"  {problem}")
else:
    print("нарушений правил оформления не найдено")
